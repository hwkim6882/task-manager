#!/usr/bin/env python3
"""
폴더 모니터링 → 문서에서 할 일 자동 추출 → Firestore 저장
"""
import os
import sys
import time
import json
import hashlib
import logging
from pathlib import Path
from datetime import datetime

from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
import firebase_admin
from firebase_admin import credentials, firestore
import anthropic

# ===== 설정 =====
WATCH_FOLDER = "/Users/macminihwkim/Library/CloudStorage/SynologyDrive-SynologyDrive/김형우 노트/[통화 회의 음성 취합] 할일 추출"
FIREBASE_KEY = "/Users/macminihwkim/task-manager/task-manager-4759b-firebase-adminsdk-fbsvc-3052457cde.json"
USER_UID = "B06hNRMchKbbN9YaT5yU4x5jin82"
PROCESSED_LOG = "/Users/macminihwkim/task-manager/.processed_files.json"
SUPPORTED_EXTS = {'.txt', '.md', '.docx', '.pdf', '.json', '.csv', '.xlsx'}
CATEGORIES = ['로펌업무', '스타트업', '마케팅', '개인', '쇼핑']

# 로깅
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s [%(levelname)s] %(message)s',
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler('/Users/macminihwkim/task-manager/watcher.log', encoding='utf-8')
    ]
)
log = logging.getLogger(__name__)

# ===== 초기화 =====
# Claude API
claude_api_key = os.environ.get('ANTHROPIC_API_KEY', '')
if not claude_api_key:
    # 설정 파일에서 읽기
    key_file = Path('/Users/macminihwkim/task-manager/.claude_api_key')
    if key_file.exists():
        claude_api_key = key_file.read_text().strip()
    else:
        log.error("ANTHROPIC_API_KEY 환경변수 또는 .claude_api_key 파일이 필요합니다.")
        log.error(f"echo 'sk-ant-...' > {key_file}")
        sys.exit(1)

client = anthropic.Anthropic(api_key=claude_api_key)

# Firebase
cred = credentials.Certificate(FIREBASE_KEY)
firebase_admin.initialize_app(cred)
db = firestore.client()

# 처리 완료 파일 목록
def load_processed():
    if os.path.exists(PROCESSED_LOG):
        with open(PROCESSED_LOG, 'r') as f:
            return json.load(f)
    return {}

def save_processed(processed):
    with open(PROCESSED_LOG, 'w') as f:
        json.dump(processed, f, ensure_ascii=False, indent=2)

def file_hash(path):
    h = hashlib.md5()
    with open(path, 'rb') as f:
        h.update(f.read())
    return h.hexdigest()

# ===== 텍스트 추출 =====
def extract_text(filepath):
    ext = Path(filepath).suffix.lower()
    try:
        if ext in ('.txt', '.md'):
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()

        elif ext == '.docx':
            from docx import Document
            doc = Document(filepath)
            return '\n'.join(p.text for p in doc.paragraphs)

        elif ext == '.pdf':
            from PyPDF2 import PdfReader
            reader = PdfReader(filepath)
            return '\n'.join(page.extract_text() or '' for page in reader.pages)

        elif ext == '.json':
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                data = json.load(f)
            # 재귀적으로 텍스트 추출
            def extract_json_text(obj):
                if isinstance(obj, str):
                    return obj
                if isinstance(obj, list):
                    return '\n'.join(extract_json_text(x) for x in obj)
                if isinstance(obj, dict):
                    return '\n'.join(f"{k}: {extract_json_text(v)}" for k, v in obj.items())
                return str(obj)
            return extract_json_text(data)

        elif ext == '.csv':
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()

        elif ext == '.xlsx':
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            text = ''
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    text += ' | '.join(str(c or '') for c in row) + '\n'
            return text

    except Exception as e:
        log.error(f"텍스트 추출 실패 [{filepath}]: {e}")
    return ''

# ===== Claude AI 할일 추출 =====
def extract_tasks_with_ai(text, source_name):
    if not text.strip() or len(text.strip()) < 10:
        return []

    categories_str = ', '.join(CATEGORIES)
    prompt = f"""당신은 법률사무소 변호사의 업무 비서입니다.
아래 텍스트는 통화 녹음 요약, 회의록, 메모 등에서 추출한 내용입니다.
출처: {source_name}

이 텍스트를 분석해서 **내가 해야 할 일(Action Items)**만 추출해주세요.

추출 기준:
- 명시적 요청/지시 사항
- 마감일이 있는 항목
- 회신/답변이 필요한 사항
- 후속 조치가 필요한 내용
- 확인/검토가 필요한 사항

추출하지 않을 것:
- 단순 정보 전달
- 이미 완료된 사항
- 상대방이 할 일

할 일이 없으면 빈 배열 []을 반환하세요.

각 할 일을 JSON 배열로 반환:
- title: 업무 제목 (구체적이고 간결하게, 관련 인물/회사명 포함)
- quad: 사분면 (q1=긴급+중요, q2=중요+긴급아님, q3=긴급+중요아님, q4=둘다아님)
- cat: 카테고리 ({categories_str} 중 하나)
- time: 실제 예상 소요시간(분) — 현실적으로 평가
- impact: 임팩트 1-5
- note: 출처 "{source_name}"

JSON 배열만 반환하세요.

텍스트:
{text[:6000]}"""

    try:
        response = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=4096,
            messages=[{"role": "user", "content": prompt}]
        )
        result_text = response.content[0].text
        # JSON 파싱
        import re
        match = re.search(r'```(?:json)?\s*([\s\S]*?)```', result_text)
        json_str = match.group(1) if match else result_text
        return json.loads(json_str.strip())
    except Exception as e:
        log.error(f"AI 할일 추출 실패: {e}")
        return []

# ===== Firestore 저장 =====
def save_tasks_to_firestore(new_tasks):
    if not new_tasks:
        return 0

    doc_ref = db.collection('users').document(USER_UID)
    doc = doc_ref.get()

    existing_tasks = []
    if doc.exists:
        data = doc.to_dict()
        existing_tasks = data.get('tasks', [])

    # 중복 체크
    existing_titles = {t.get('title', '').strip().lower().replace(' ', '') for t in existing_tasks}
    added = 0

    for task in new_tasks:
        norm_title = task.get('title', '').strip().lower().replace(' ', '')
        if norm_title in existing_titles:
            log.info(f"  중복 건너뜀: {task['title']}")
            continue

        existing_tasks.append({
            'id': time.time() + added * 0.001,
            'title': task.get('title', ''),
            'quadrant': task.get('quad', 'q3'),
            'category': task.get('cat', '로펌업무'),
            'timeMin': task.get('time', 30),
            'impact': task.get('impact', 3),
            'note': task.get('note', ''),
            'done': False,
            'createdAt': datetime.now().isoformat(),
            'doneAt': None,
            'autoExtracted': True,
        })
        existing_titles.add(norm_title)
        added += 1

    if added > 0:
        doc_ref.update({'tasks': existing_tasks})

    return added

# ===== 파일 처리 =====
def process_file(filepath):
    filepath = str(filepath)
    ext = Path(filepath).suffix.lower()

    if ext not in SUPPORTED_EXTS:
        return
    if '.DS_Store' in filepath:
        return

    processed = load_processed()
    fhash = file_hash(filepath)

    if filepath in processed and processed[filepath] == fhash:
        return  # 이미 처리됨

    # 같은 이름(확장자만 다른) 파일이 이미 처리되었으면 건너뜀
    filename = Path(filepath).name
    if filename.endswith('.summary.txt'):
        stem = filename[:-len('.summary.txt')]
    else:
        stem = Path(filepath).stem
    parent_dir = str(Path(filepath).parent)
    for ppath in processed:
        pname = Path(ppath).name
        pstem = pname[:-len('.summary.txt')] if pname.endswith('.summary.txt') else Path(ppath).stem
        if pstem == stem and str(Path(ppath).parent) == parent_dir and ppath != filepath:
            log.info(f"  같은 이름 파일 이미 처리됨, 건너뜀: {filename}")
            processed[filepath] = fhash
            save_processed(processed)
            return

    parent = Path(filepath).parent.name
    source = f"{parent}/{filename}"

    log.info(f"새 파일 감지: {source}")

    # 텍스트 추출
    text = extract_text(filepath)
    if not text.strip():
        log.info(f"  텍스트 없음, 건너뜀")
        processed[filepath] = fhash
        save_processed(processed)
        return

    log.info(f"  텍스트 추출 완료 ({len(text)}자)")

    # AI 할일 추출
    tasks = extract_tasks_with_ai(text, source)
    log.info(f"  AI 추출: {len(tasks)}건")

    if tasks:
        added = save_tasks_to_firestore(tasks)
        log.info(f"  Firestore 저장: {added}건 추가")
        for t in tasks:
            log.info(f"    → [{t.get('cat','')}] {t.get('title','')}")

    # 처리 완료 기록
    processed[filepath] = fhash
    save_processed(processed)

# ===== 폴더 감시 =====
class TaskFileHandler(FileSystemEventHandler):
    def __init__(self):
        self.debounce = {}

    def on_created(self, event):
        if event.is_directory:
            return
        self._handle(event.src_path)

    def on_modified(self, event):
        if event.is_directory:
            return
        self._handle(event.src_path)

    def _handle(self, path):
        # 디바운스 (같은 파일 2초 이내 재처리 방지)
        now = time.time()
        if path in self.debounce and now - self.debounce[path] < 2:
            return
        self.debounce[path] = now

        # 파일 쓰기 완료 대기
        time.sleep(1)
        try:
            process_file(path)
        except Exception as e:
            log.error(f"파일 처리 오류 [{path}]: {e}")

# ===== 메인 =====
def group_files_by_stem(folder):
    """같은 이름(확장자만 다른) 파일을 그룹화하고, 우선순위로 1개만 선택"""
    EXT_PRIORITY = ['.summary.txt', '.md', '.txt', '.docx', '.pdf', '.json', '.csv', '.xlsx']
    groups = {}

    for root, dirs, files in os.walk(folder):
        for fname in files:
            filepath = os.path.join(root, fname)
            ext = Path(filepath).suffix.lower()
            if ext not in SUPPORTED_EXTS:
                continue
            if '.DS_Store' in filepath:
                continue

            # .summary.txt 특수 처리
            if fname.endswith('.summary.txt'):
                stem = os.path.join(root, fname[:-len('.summary.txt')])
            else:
                stem = os.path.join(root, Path(filepath).stem)

            if stem not in groups:
                groups[stem] = []
            groups[stem].append(filepath)

    # 각 그룹에서 우선순위가 높은 파일 1개만 선택
    selected = []
    for stem, filepaths in groups.items():
        def priority(fp):
            for i, ext in enumerate(EXT_PRIORITY):
                if fp.endswith(ext):
                    return i
            return len(EXT_PRIORITY)
        filepaths.sort(key=priority)
        selected.append(filepaths[0])

    return selected

def initial_scan():
    """기존 파일 중 미처리 항목 스캔 (같은 이름 파일은 1개만)"""
    log.info("기존 파일 스캔 시작...")
    processed = load_processed()
    count = 0

    selected_files = group_files_by_stem(WATCH_FOLDER)
    for filepath in selected_files:
        fhash = file_hash(filepath)
        if filepath in processed and processed[filepath] == fhash:
            continue

        count += 1
        process_file(filepath)
        time.sleep(0.5)  # API 속도 제한 방지

    log.info(f"초기 스캔 완료: {count}건 처리")

if __name__ == '__main__':
    log.info("=" * 50)
    log.info("Task Folder Watcher 시작")
    log.info(f"감시 폴더: {WATCH_FOLDER}")
    log.info("=" * 50)

    # 초기 스캔 (--scan 옵션)
    if '--scan' in sys.argv:
        initial_scan()

    # 폴더 감시 시작
    handler = TaskFileHandler()
    observer = Observer()
    observer.schedule(handler, WATCH_FOLDER, recursive=True)
    observer.start()

    log.info("폴더 감시 중... (Ctrl+C로 종료)")

    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        observer.stop()
        log.info("종료")
    observer.join()
